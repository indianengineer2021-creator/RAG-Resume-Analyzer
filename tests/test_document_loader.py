from pathlib import Path

from docx import Document

from src.document_loader import load_resume_documents


def test_load_resume_documents_from_docx(tmp_path):
    file_path = tmp_path / "resume.docx"
    doc = Document()
    doc.add_paragraph("Senior Data Engineer")
    doc.add_paragraph("Python, SQL, Azure")
    doc.save(file_path)

    docs = load_resume_documents(file_path)

    assert len(docs) == 1
    assert docs[0].metadata["document_type"] == "resume"
    assert "Python" in docs[0].page_content


def test_load_resume_documents_from_pdf(monkeypatch, tmp_path):
    file_path = tmp_path / "resume.pdf"
    file_path.write_bytes(b"%PDF-1.4")

    def fake_extract_pdf_text(_path):
        return "Machine Learning Engineer\nPython, TensorFlow, AWS"

    monkeypatch.setattr("src.document_loader.extract_pdf_text", fake_extract_pdf_text)

    docs = load_resume_documents(file_path)

    assert len(docs) == 1
    assert docs[0].metadata["source"] == "resume.pdf"
    assert "TensorFlow" in docs[0].page_content
