import logging
from pathlib import Path
from typing import List, Union

from docx import Document as DocxDocument
from langchain_core.documents import Document
from pypdf import PdfReader

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_pdf_text(file_path: Union[str, Path]) -> str:
    """Read a PDF file and return the combined text from all pages."""
    file_path = Path(file_path)
    text_chunks: List[str] = []
    reader = PdfReader(str(file_path))

    if len(reader.pages) == 0:
        raise ValueError(f"No readable pages found in PDF: {file_path.name}")

    for page_number, page in enumerate(reader.pages, start=1):
        page_text = page.extract_text() or ""
        if page_text.strip():
            text_chunks.append(f"Page {page_number}\n{page_text}")

    full_text = "\n\n".join(text_chunks).strip()
    if not full_text:
        raise ValueError(f"PDF extraction returned empty text for: {file_path.name}")
    return full_text


def extract_docx_text(file_path: Union[str, Path]) -> str:
    """Read a DOCX file and return the document text."""
    file_path = Path(file_path)
    doc = DocxDocument(str(file_path))
    paragraphs = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text and paragraph.text.strip()]
    full_text = "\n".join(paragraphs).strip()

    if not full_text:
        raise ValueError(f"DOCX extraction returned empty text for: {file_path.name}")
    return full_text


def load_resume_documents(file_path: Union[str, Path]) -> List[Document]:
    """Load a PDF or DOCX resume and return LangChain documents with metadata."""
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"Resume file not found: {file_path}")

    suffix = file_path.suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {suffix}. Please upload a PDF or DOCX resume.")

    try:
        if suffix == ".pdf":
            text = extract_pdf_text(file_path)
        else:
            text = extract_docx_text(file_path)
    except Exception as exc:  # pragma: no cover - defensive logging path
        logger.exception("Resume parsing failed for %s", file_path)
        raise ValueError(f"Failed to extract text from resume: {file_path.name}. Error: {exc}") from exc

    if not text or not text.strip():
        raise ValueError(f"The uploaded resume is empty: {file_path.name}")

    metadata = {"source": file_path.name, "document_type": "resume"}
    return [Document(page_content=text, metadata=metadata)]


def load_jd_document(file_path: Union[str, Path]) -> str:
    """Load a JD from a text-like file or PDF/DOCX and return plain text."""
    file_path = Path(file_path)
    if file_path.suffix.lower() == ".pdf":
        return extract_pdf_text(file_path)
    if file_path.suffix.lower() == ".docx":
        return extract_docx_text(file_path)
    if file_path.suffix.lower() in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"Unsupported JD format: {file_path.suffix}. Please upload a text, PDF, or DOCX job description.")
